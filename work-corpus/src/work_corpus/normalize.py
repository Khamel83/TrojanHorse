from __future__ import annotations

import csv
from email import policy
from email.message import Message
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import io
import json
import posixpath
import re
import sqlite3
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple

from .config import Config, DEFAULT_SKIP_CLASSIFICATIONS
from .db import (
    mark_noncurrent_normalized_documents_retained,
    record_evidence,
    record_review_item,
)
from .entities import record_scope_proposal
from .onenote import convert_one
from .tasks import extract_task_proposals
from .util import (
    atomic_write_json,
    atomic_write_text,
    ensure_dir,
    effective_extension,
    html_to_text,
    now_iso,
    provenance_header,
    read_text_guess,
    scrub_derived_text,
    sha256_text,
    slugify,
    stable_evidence_id,
    stable_id,
    vtt_or_srt_to_markdown,
)


# Bumped after broadening pointer-only Capacities metadata detection. Existing
# derived files must be rebuilt instead of being treated as unchanged.
PARSER_VERSION = "v2"
EMAIL_PARSER_VERSION = "v1"
TIMESTAMP_LINE = re.compile(
    r"^\s*(?:\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{3}\s*-->"
    r"\s*(?:\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{3}"
)


@dataclass(frozen=True)
class ExtractedPart:
    locator: str
    text: str


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    parser: str
    parser_version: str = PARSER_VERSION
    parts: Tuple[ExtractedPart, ...] = ()
    status: str = "normalized"
    metadata: Mapping[str, Any] = field(default_factory=dict)


class ExtractionBlocked(RuntimeError):
    """Raised when a local adapter is intentionally unavailable."""


def _output_path(config: Config, source_system: str, source_version_id: str) -> Path:
    output = (
        config.corpus_dir
        / "normalized"
        / slugify(str(source_system))
        / f"{source_version_id}.md"
    )
    config.assert_derived_path(output)
    return output


def _has_explicit_source_root(row: sqlite3.Row) -> bool:
    """Require the immutable inventory classification before semantic parsing."""
    try:
        metadata = json.loads(row["metadata_json"] or "{}")
    except (TypeError, json.JSONDecodeError):
        return False
    return metadata.get("root_match_kind") in {"source", "archive"}


def _parse_csv(path: Path, delimiter: str, max_rows: int) -> str:
    text = read_text_guess(path)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    output: List[str] = []
    for index, row in enumerate(reader):
        if index >= max_rows:
            output.append("")
            output.append(f"[TRUNCATED AFTER {max_rows} ROWS]")
            break
        output.append("\t".join(cell.replace("\t", " ").replace("\r", " ").replace("\n", " ") for cell in row))
    return "\n".join(output) + "\n"


def _parse_json(path: Path) -> str:
    text = read_text_guess(path)
    if path.suffix.lower() == ".jsonl":
        lines = []
        for number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                lines.append(json.dumps(obj, ensure_ascii=False, sort_keys=True))
            except json.JSONDecodeError:
                lines.append(f"[UNPARSEABLE JSONL LINE {number}] {line}")
        return "\n".join(lines) + "\n"
    payload = json.loads(text)
    return json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True) + "\n"


def _parse_rtf(path: Path) -> str:
    text = read_text_guess(path)
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip() + "\n"


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise RuntimeError("optional dependency pypdf is not installed") from exc
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            page_text = f"[PAGE {index} EXTRACTION ERROR: {exc}]"
        pages.append(f"## Page {index}\n\n{page_text.strip()}\n")
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("optional dependency python-docx is not installed") from exc
    document = Document(str(path))
    output: List[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style.startswith("heading"):
                digits = re.findall(r"\d+", style)
                level = min(int(digits[0]), 6) if digits else 2
                output.append("#" * level + " " + paragraph.text.strip())
            else:
                output.append(paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        output.append(f"\n## Table {table_index}\n")
        for row in table.rows:
            output.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
    return "\n\n".join(output).strip() + "\n"


def _parse_pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise RuntimeError("optional dependency python-pptx is not installed") from exc
    prs = Presentation(str(path))
    output: List[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        output.append(f"# Slide {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    output.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    output.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
        notes_slide = getattr(slide, "notes_slide", None)
        if notes_slide:
            note_parts = []
            for shape in notes_slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value and value.lower() not in {"slide image", "text placeholder"}:
                        note_parts.append(value)
            if note_parts:
                output.append("## Speaker notes")
                output.extend(note_parts)
        output.append("")
    return "\n\n".join(output).strip() + "\n"


def _parse_xlsx(path: Path, max_rows: int) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError("optional dependency openpyxl is not installed") from exc
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    output: List[str] = []
    remaining = max_rows
    for sheet in workbook.worksheets:
        output.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if remaining <= 0:
                output.append(f"[WORKBOOK TRUNCATED AFTER {max_rows} TOTAL ROWS]")
                return "\n".join(output) + "\n"
            output.append("\t".join("" if value is None else str(value).replace("\n", " ") for value in row))
            remaining -= 1
        output.append("")
    return "\n".join(output)


def _sqlite_cell(value: Any, max_chars: int = 10000) -> str:
    if value is None:
        return ""
    if isinstance(value, (bytes, bytearray, memoryview)):
        return f"[BLOB {len(value)} bytes]"
    text = str(value).replace("\r", " ").replace("\n", "\\n").replace("\t", " ")
    if len(text) > max_chars:
        return text[:max_chars] + f"… [TRUNCATED {len(text) - max_chars} CHARS]"
    return text


def _parse_sqlite(path: Path, max_rows: int) -> str:
    output = [f"# SQLite read-only extraction: {path.name}", ""]
    uri = f"file:{path.as_posix()}?mode=ro&immutable=1"
    con = sqlite3.connect(uri, uri=True)
    remaining = max_rows
    try:
        tables = con.execute(
            "SELECT name, type, sql FROM sqlite_master "
            "WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY type, name"
        ).fetchall()
        for name, obj_type, sql in tables:
            output.append(f"## {obj_type.title()}: {name}")
            output.append("")
            if sql:
                output.append("```sql")
                output.append(sql)
                output.append("```")
            escaped_name = name.replace('"', '""')
            count = None
            if obj_type == "table":
                try:
                    count = con.execute(f'SELECT COUNT(*) FROM "{escaped_name}"').fetchone()[0]
                    output.append(f"Rows: {count}")
                except Exception as exc:
                    output.append(f"Rows: [COUNT ERROR: {exc}]")
            try:
                cols = con.execute(f'PRAGMA table_info("{escaped_name}")').fetchall()
                column_names = [str(col[1]) for col in cols]
                if column_names:
                    output.append("")
                    output.append("Columns:")
                    for col in cols:
                        output.append(f"- {col[1]} ({col[2] or 'untyped'})")
            except Exception as exc:
                column_names = []
                output.append(f"- [COLUMN ERROR: {exc}]")

            if remaining <= 0:
                output.append("")
                output.append(f"[ROW EXTRACTION STOPPED AFTER {max_rows} TOTAL DATABASE ROWS]")
                continue

            # Read views and tables conservatively. Cell bodies are capped and BLOBs
            # are represented by length rather than copied into normalized text.
            try:
                cursor = con.execute(f'SELECT * FROM "{escaped_name}" LIMIT ?', (remaining,))
                headers = [description[0] for description in cursor.description or []]
                rows = cursor.fetchall()
                if headers:
                    output.append("")
                    output.append("### Extracted rows")
                    output.append("")
                    output.append("```tsv")
                    output.append("\t".join(_sqlite_cell(value) for value in headers))
                    for row in rows:
                        output.append("\t".join(_sqlite_cell(value) for value in row))
                    output.append("```")
                    remaining -= len(rows)
                    if count is not None and count > len(rows):
                        output.append(
                            f"[TABLE TRUNCATED: extracted {len(rows)} of {count} rows; "
                            f"database-wide limit is {max_rows}]"
                        )
            except Exception as exc:
                output.append(f"[ROW EXTRACTION ERROR: {exc}]")
            output.append("")
    finally:
        con.close()
    return "\n".join(output)


def _parse_zip(path: Path, max_entries: int) -> str:
    output = [f"# Archive inventory: {path.name}", ""]
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        output.append(f"Entries: {len(infos)}")
        output.append("")
        for index, info in enumerate(infos):
            if index >= max_entries:
                output.append(f"[TRUNCATED AFTER {max_entries} ENTRIES]")
                break
            encrypted = bool(info.flag_bits & 0x1)
            output.append(
                f"- `{info.filename}` — compressed {info.compress_size} bytes; "
                f"uncompressed {info.file_size} bytes; encrypted={str(encrypted).lower()}"
            )
    return "\n".join(output) + "\n"


def _render_csv_row(row: Iterable[str]) -> str:
    return "\t".join(
        cell.replace("\t", " ").replace("\r", " ").replace("\n", " ")
        for cell in row
    )


def _csv_parts(path: Path, delimiter: str, max_rows: int) -> List[ExtractedPart]:
    text = read_text_guess(path)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: List[List[str]] = []
    truncated = False
    for index, row in enumerate(reader):
        if index >= max_rows:
            truncated = True
            break
        rows.append(row)
    if not rows:
        parts = [ExtractedPart("document", "")]
    else:
        parts = [ExtractedPart("header", _render_csv_row(rows[0]))]
        parts.extend(
            ExtractedPart(f"row:{index}", _render_csv_row(row))
            for index, row in enumerate(rows[1:], start=1)
        )
    if truncated:
        parts.append(
            ExtractedPart("truncated", f"[TRUNCATED AFTER {max_rows} ROWS]")
        )
    return parts


def _html_parts(path: Path, max_text: int) -> List[ExtractedPart]:
    return [
        ExtractedPart(
            "document",
            html_to_text(read_text_guess(path, max_bytes=max_text)) + "\n",
        )
    ]


def _text_parts(path: Path, max_text: int) -> List[ExtractedPart]:
    return [ExtractedPart("document", read_text_guess(path, max_bytes=max_text))]


def _email_payload(part: Message) -> str:
    try:
        value = part.get_content()
        if isinstance(value, str):
            return value
    except Exception:
        pass
    raw = part.get_payload(decode=True)
    if raw is None:
        value = part.get_payload()
        return value if isinstance(value, str) else ""
    charset = part.get_content_charset() or "utf-8"
    try:
        return raw.decode(charset, errors="replace")
    except LookupError:
        return raw.decode("utf-8", errors="replace")


def _email_date(value: str) -> Optional[str]:
    if not value:
        return None
    try:
        parsed = parsedate_to_datetime(value)
    except (TypeError, ValueError, IndexError, OverflowError):
        return None
    return parsed.date().isoformat() if parsed else None


def _email_parts(path: Path, max_bytes: int) -> ExtractionResult:
    raw = path.read_bytes()
    if len(raw) > max_bytes:
        raise RuntimeError(
            f"email source exceeds configured limit ({len(raw)} > {max_bytes} bytes)"
        )
    message = BytesParser(policy=policy.default).parsebytes(raw)
    plain_parts: List[str] = []
    html_parts: List[str] = []
    attachments: List[Dict[str, Any]] = []
    for part in message.walk():
        if part.is_multipart():
            continue
        filename = part.get_filename()
        disposition = (part.get_content_disposition() or "").casefold()
        content_type = part.get_content_type()
        if disposition == "attachment" or filename:
            payload = part.get_payload(decode=True)
            attachments.append(
                {
                    "name": str(filename or ""),
                    "content_type": content_type,
                    "size_bytes": len(payload) if payload is not None else 0,
                }
            )
            continue
        if content_type == "text/plain":
            plain_parts.append(_email_payload(part))
        elif content_type == "text/html":
            html_parts.append(_email_payload(part))
    body = "\n\n".join(part.strip() for part in plain_parts if part.strip())
    if not body and html_parts:
        body = html_to_text("\n".join(html_parts))
    body = body[:max_bytes]

    date_header = str(message.get("Date") or "").strip()
    event_date = _email_date(date_header)
    header_lines = [
        f"Subject: {str(message.get('Subject') or '').strip()}",
        f"From: {str(message.get('From') or '').strip()}",
        f"To: {'; '.join(str(value).strip() for value in message.get_all('To', []))}",
        f"Cc: {'; '.join(str(value).strip() for value in message.get_all('Cc', []))}",
        f"Date: {date_header}",
        f"Message-ID: {str(message.get('Message-ID') or '').strip()}",
    ]
    attachment_lines = [
        f"- {item['name'] or '(unnamed)'} — {item['content_type']}; "
        f"{item['size_bytes']} bytes"
        for item in attachments
    ]
    return _result(
        "email_eml:v1",
        [
            ExtractedPart("headers", "\n".join(header_lines)),
            ExtractedPart("body", body),
            ExtractedPart(
                "attachments",
                "\n".join(attachment_lines)
                if attachment_lines
                else "Attachments: none",
            ),
        ],
        parser_version=EMAIL_PARSER_VERSION,
        metadata={
            "event_date": event_date or "",
            "source_date_basis": "email_header_date" if event_date else "",
            "attachment_count": len(attachments),
        },
    )


def _caption_parts(path: Path, max_text: int) -> List[ExtractedPart]:
    lines = read_text_guess(path, max_bytes=max_text).replace("\ufeff", "").splitlines()
    parts: List[ExtractedPart] = []
    current_stamp = ""
    current_text: List[str] = []

    def flush() -> None:
        nonlocal current_stamp, current_text
        body = " ".join(
            part.strip() for part in current_text if part.strip()
        ).strip()
        body = re.sub(r"<[^>]+>", "", body)
        if body:
            locator = (
                f"timestamp:{current_stamp.split('-->', 1)[0].strip()}"
                if current_stamp
                else f"cue:{len(parts) + 1}"
            )
            parts.append(
                ExtractedPart(
                    locator,
                    f"**{current_stamp}**\n{body}" if current_stamp else body,
                )
            )
        current_stamp = ""
        current_text = []

    for raw in lines:
        line = raw.strip()
        if not line or line == "WEBVTT" or line.startswith("NOTE"):
            flush()
            continue
        if line.isdigit() and not current_text:
            continue
        if TIMESTAMP_LINE.match(line):
            flush()
            current_stamp = line
            continue
        if line.startswith(("Kind:", "Language:")):
            continue
        current_text.append(line)
    flush()
    return parts or [ExtractedPart("document", "")]


def _pdf_parts(path: Path) -> List[ExtractedPart]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise RuntimeError("optional dependency pypdf is not installed") from exc
    reader = PdfReader(str(path))
    parts: List[ExtractedPart] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            page_text = f"[PAGE {index} EXTRACTION ERROR: {exc}]"
        parts.append(
            ExtractedPart(
                f"page:{index}",
                f"## Page {index}\n\n{page_text.strip()}".rstrip(),
            )
        )
    return parts or [ExtractedPart("document", "")]


def _docx_parts(path: Path) -> List[ExtractedPart]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "optional dependency python-docx is not installed"
        ) from exc
    document = Document(str(path))
    parts: List[ExtractedPart] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            digits = re.findall(r"\d+", style)
            level = min(int(digits[0]), 6) if digits else 2
            text = "#" * level + " " + text
        parts.append(ExtractedPart(f"paragraph:{index}", text))
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            parts.append(
                ExtractedPart(
                    f"table:{table_index}:row:{row_index}",
                    _render_csv_row(cell.text for cell in row.cells),
                )
            )
    return parts or [ExtractedPart("document", "")]


def _pptx_parts(path: Path) -> List[ExtractedPart]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "optional dependency python-pptx is not installed"
        ) from exc
    presentation = Presentation(str(path))
    parts: List[ExtractedPart] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        output = [f"# Slide {slide_index}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    output.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    output.append(_render_csv_row(cell.text for cell in row.cells))
        notes_slide = getattr(slide, "notes_slide", None)
        if notes_slide:
            notes = []
            for shape in notes_slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value and value.casefold() not in {
                        "slide image",
                        "text placeholder",
                    }:
                        notes.append(value)
            if notes:
                output.append("## Speaker notes")
                output.extend(notes)
        parts.append(ExtractedPart(f"slide:{slide_index}", "\n\n".join(output)))
    return parts or [ExtractedPart("document", "")]


def _xlsx_parts(path: Path, max_rows: int) -> List[ExtractedPart]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "optional dependency openpyxl is not installed"
        ) from exc
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: List[ExtractedPart] = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > max_rows:
                parts.append(
                    ExtractedPart(
                        f"sheet:{sheet.title}:truncated",
                        f"[SHEET TRUNCATED AFTER {max_rows} ROWS]",
                    )
                )
                break
            parts.append(
                ExtractedPart(
                    f"sheet:{sheet.title}:row:{row_index}",
                    _render_csv_row(
                        "" if value is None else str(value) for value in row
                    ),
                )
            )
    return parts or [ExtractedPart("document", "")]


def _zip_parts(path: Path, max_entries: int) -> List[ExtractedPart]:
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
    parts: List[ExtractedPart] = []
    for index, info in enumerate(infos[:max_entries], start=1):
        encrypted = bool(info.flag_bits & 0x1)
        parts.append(
            ExtractedPart(
                f"archive:member:{index}",
                (
                    f"[{info.filename}] - compressed {info.compress_size} bytes; "
                    f"uncompressed {info.file_size} bytes; "
                    f"encrypted={str(encrypted).lower()}"
                ),
            )
        )
    if len(infos) > max_entries:
        parts.append(
            ExtractedPart(
                "archive:truncated",
                f"[TRUNCATED AFTER {max_entries} ENTRIES]",
            )
        )
    return parts or [ExtractedPart("archive", "")]


def _result(
    parser: str,
    parts: Iterable[ExtractedPart],
    *,
    status: str = "normalized",
    metadata: Optional[Mapping[str, Any]] = None,
    parser_version: str = PARSER_VERSION,
) -> ExtractionResult:
    locator_counts: Dict[str, int] = {}
    cleaned_parts: List[ExtractedPart] = []
    for part in parts:
        base_locator = part.locator or "document"
        count = locator_counts.get(base_locator, 0) + 1
        locator_counts[base_locator] = count
        locator = base_locator if count == 1 else f"{base_locator}#part:{count}"
        while any(existing.locator == locator for existing in cleaned_parts):
            count += 1
            locator_counts[base_locator] = count
            locator = f"{base_locator}#part:{count}"
        cleaned_parts.append(
            ExtractedPart(locator, scrub_derived_text(part.text))
        )
    cleaned = tuple(cleaned_parts)
    body = "\n\n".join(part.text.rstrip() for part in cleaned if part.text.strip())
    return ExtractionResult(
        text=body + ("\n" if body else ""),
        parser=parser,
        parser_version=parser_version,
        parts=cleaned or (ExtractedPart("document", ""),),
        status=status,
        metadata=dict(metadata or {}),
    )


def _pointer_only(text: str) -> bool:
    pointer_type = re.search(
        r"(?im)^\s*(?:type|kind|content[_ -]?type)\s*:\s*"
        r"['\"]?(?:file|image|pdf)['\"]?\s*$",
        text,
    )
    pointer_field = re.search(
        r"(?im)^\s*(?:url|media|download[_ -]?url|signed[_ -]?url)\s*:",
        text,
    )
    local_payload = re.search(
        r"(?im)^\s*(?:local[_ -]?(?:path|file)|attachment[_ -]?path|"
        r"payload[_ -]?path)\s*:\s*\S+",
        text,
    )
    return bool(pointer_type and pointer_field and not local_payload)


def _notion_or_capacities_parts(
    path: Path,
    config: Config,
    source_system: str,
    kind: str,
    extension: str,
) -> ExtractionResult:
    max_text = int(config.get("normalization", "max_text_file_mb", 200)) * 1024 * 1024
    max_rows = int(config.get("normalization", "max_table_rows", 100000))
    ext = (
        effective_extension(path)
        if re.search(r"\.[A-Za-z0-9]+\s+\(\d+\)$", path.name)
        else extension.lower()
    )
    if source_system == "capacities" and ext in {".md", ".markdown"}:
        raw = read_text_guess(path, max_bytes=max_text)
        pointer = _pointer_only(raw)
        if pointer:
            raw = "payload_status: missing\n\n" + raw
        return _result(
            "capacities_markdown:v1",
            [ExtractedPart("document", raw)],
            metadata={"payload_status": "missing" if pointer else "present"},
        )
    if source_system == "capacities" and ext == ".csv":
        raw = read_text_guess(path, max_bytes=max_text)
        parts = _csv_parts(path, ",", max_rows)
        pointer = bool(
            re.search(
                r"(?i)\b(?:file|image|pdf)\b",
                raw,
            )
            and re.search(r"(?i)\b(?:url|path|payload)\b", raw)
        )
        if pointer:
            parts = [
                ExtractedPart("payload_status", "payload_status: missing"),
                *parts,
            ]
        return _result(
            "capacities_csv:v1",
            parts,
            metadata={"payload_status": "missing" if pointer else "present"},
        )
    if source_system == "capacities" and ext in {".json", ".jsonl"}:
        return _result(
            "capacities_json:v1",
            [ExtractedPart("document", _parse_json(path))],
        )
    if source_system == "capacities" and ext in {".html", ".htm"}:
        return _result("capacities_html:v1", _html_parts(path, max_text))
    if source_system == "capacities" and ext == ".txt":
        return _result("capacities_text:v1", _text_parts(path, max_text))
    if source_system == "notion" and ext in {".html", ".htm"}:
        return _result("notion_html:v1", _html_parts(path, max_text))
    if source_system == "notion" and ext == ".csv":
        return _result(
            "notion_database_csv:v1",
            _csv_parts(path, ",", max_rows),
        )
    if source_system == "notion" and ext == ".txt":
        return _result("notion_transcript_txt:v1", _text_parts(path, max_text))
    if source_system == "notion" and ext in {".md", ".markdown"}:
        return _result("notion_markdown:v1", _text_parts(path, max_text))
    if source_system == "notion":
        formal = _formal_parts(path, ext, max_rows)
        return _result(
            f"notion_attachment_{slugify(ext.lstrip('.'))}:v1",
            formal,
            metadata={"evidence_unit": "local_attachment"},
        )
    return _result(
        f"{source_system}_document:v1",
        _text_parts(path, max_text),
    )


def _formal_parts(
    path: Path,
    extension: str,
    max_rows: int,
) -> List[ExtractedPart]:
    ext = extension.lower()
    if ext == ".pdf":
        return _pdf_parts(path)
    if ext == ".docx":
        return _docx_parts(path)
    if ext == ".pptx":
        return _pptx_parts(path)
    if ext == ".xlsx":
        return _xlsx_parts(path, max_rows)
    if ext == ".csv":
        return _csv_parts(path, ",", max_rows)
    if ext == ".tsv":
        return _csv_parts(path, "\t", max_rows)
    raise RuntimeError(f"unsupported normalization format: {ext or '[no extension]'}")


def _onenote_parts(text: str) -> List[ExtractedPart]:
    matches = list(
        re.finditer(
            r"(?im)^(?P<heading>#+\s*(?:section|page)\b[^\n]*)$",
            text,
        )
    )
    if not matches:
        return [ExtractedPart("document", text)]
    parts: List[ExtractedPart] = []
    for index, match in enumerate(matches):
        heading = match.group("heading").strip().lstrip("#").strip()
        if ":" in heading:
            kind, value = heading.split(":", 1)
        elif " " in heading:
            kind, value = heading.split(None, 1)
        else:
            kind, value = heading, str(index + 1)
        locator = f"{kind.casefold()}:{value.strip()}"
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        parts.append(ExtractedPart(locator, text[match.start():end].strip()))
    return parts


def extract_source(
    path: Path,
    kind: str,
    extension: str,
    config: Config,
    *,
    source_system: Optional[str] = None,
    relative_path: str = "",
) -> ExtractionResult:
    """Extract one immutable source version into scrubbed, locatable parts."""
    system = (source_system or "").casefold()
    ext = (
        effective_extension(path)
        if re.search(r"\.[A-Za-z0-9]+\s+\(\d+\)$", path.name)
        else extension.lower()
    )
    max_text = int(config.get("normalization", "max_text_file_mb", 200)) * 1024 * 1024
    max_rows = int(config.get("normalization", "max_table_rows", 100000))

    if ext == ".eml":
        return _email_parts(path, max_text)

    if ext in {".one", ".onepkg"}:
        converted = convert_one(config, path)
        if converted.status == "blocked":
            raise ExtractionBlocked(converted.error)
        if converted.status != "succeeded":
            raise RuntimeError(converted.error or "OneNote conversion failed")
        return _result(
            converted.parser,
            _onenote_parts(converted.text),
            metadata={
                **converted.metadata,
                "converter": converted.converter,
                "relative_path": relative_path,
            },
        )

    if system in {"capacities", "notion"}:
        return _notion_or_capacities_parts(
            path,
            config,
            system,
            kind,
            ext,
        )

    if system == "zoom":
        name = path.name.casefold()
        if kind == "metadata" or name.startswith("client_config"):
            return _result(
                "zoom_client_config:v1",
                _text_parts(path, max_text),
                status="metadata",
            )
        if ext in {".vtt", ".srt"}:
            return _result(
                f"zoom_transcript_{ext.lstrip('.')}:v1",
                _caption_parts(path, max_text),
            )
        if ext == ".txt" and (
            kind in {"transcript", "transcript_candidate"}
            or any(token in name for token in ("transcript", "caption"))
        ):
            return _result("zoom_transcript_txt:v1", _text_parts(path, max_text))
        if kind == "meeting_chat":
            return _result(
                "zoom_meeting_chat:v1",
                _text_parts(path, max_text),
                status="metadata",
            )

    if kind in {"transcript", "transcript_candidate"} and ext in {
        ".md",
        ".markdown",
        ".txt",
    }:
        return _result(
            f"transcript_{ext.lstrip('.')}:v1",
            _text_parts(path, max_text),
        )
    if ext in {".md", ".markdown", ".txt"}:
        return _result("plain_text:v1", _text_parts(path, max_text))
    if ext in {".html", ".htm"}:
        return _result("html_to_text:v1", _html_parts(path, max_text))
    if ext == ".rtf":
        return _result("rtf_basic:v1", [ExtractedPart("document", _parse_rtf(path))])
    if ext in {".csv", ".tsv"}:
        return _result(
            "csv:v1" if ext == ".csv" else "tsv:v1",
            _csv_parts(path, "," if ext == ".csv" else "\t", max_rows),
        )
    if ext in {".json", ".jsonl"}:
        return _result("json:v1", [ExtractedPart("document", _parse_json(path))])
    if ext in {".vtt", ".srt"}:
        return _result(
            f"transcript_{ext.lstrip('.')}:v1",
            _caption_parts(path, max_text),
        )
    if ext in {".pdf", ".docx", ".pptx", ".xlsx"}:
        parser_name = {
            ".pdf": "formal_pdf:v1",
            ".docx": "formal_docx:v1",
            ".pptx": "formal_pptx:v1",
            ".xlsx": "formal_xlsx:v1",
        }[ext]
        return _result(parser_name, _formal_parts(path, ext, max_rows))
    if ext in {".sqlite", ".sqlite3", ".db"}:
        return _result(
            "sqlite_read_only:v1",
            [ExtractedPart("database", _parse_sqlite(path, max_rows))],
        )
    if ext == ".zip":
        max_entries = int(
            config.get("inventory", "max_archive_entries_to_list", 100000)
        )
        return _result("zip_inventory:v1", _zip_parts(path, max_entries))

    if ext in {".olm", ".pst", ".ost"}:
        raise RuntimeError(
            "proprietary archive: retain as raw evidence; convert with an approved local tool"
        )
    if ext in {".doc", ".ppt", ".xls"}:
        raise RuntimeError("legacy Microsoft Office binary format requires conversion")
    raise RuntimeError(f"unsupported normalization format: {ext or '[no extension]'}")


def parse_source(
    path: Path,
    kind: str,
    extension: str,
    config: Config,
    *,
    source_system: Optional[str] = None,
    relative_path: str = "",
) -> Tuple[str, str]:
    result = extract_source(
        path,
        kind,
        extension,
        config,
        source_system=source_system,
        relative_path=relative_path,
    )
    return result.text, result.parser


def _row_metadata(row: sqlite3.Row) -> Dict[str, Any]:
    try:
        value = json.loads(row["metadata_json"] or "{}")
    except (TypeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _archive_directory_path(
    config: Config,
    archive_row: sqlite3.Row,
    member_path: str,
) -> Optional[str]:
    metadata = _row_metadata(archive_row)
    extracted_key = str(metadata.get("archive_extracted_root_key") or "")
    if not extracted_key:
        return None
    extracted_root = next(
        (root for root in config.source_roots if root.key == extracted_key),
        None,
    )
    if extracted_root is None:
        return None
    member = posixpath.normpath(str(member_path).replace("\\", "/"))
    if not member or member == "." or member.startswith("../"):
        return None
    root_name = posixpath.basename(extracted_root.relative_path.rstrip("/"))
    if member == root_name:
        return None
    if member.startswith(root_name + "/"):
        member = member[len(root_name) + 1 :]
    return posixpath.join(extracted_root.relative_path, member)


def build_archive_reference_map(
    config: Config,
    con: sqlite3.Connection,
) -> Dict[str, Dict[str, Any]]:
    """Map verified archive members to extracted semantic evidence once."""
    rows = con.execute(
        """
        SELECT *
        FROM source_record
        WHERE kind='archive'
          AND status='present'
        ORDER BY relative_path
        """
    ).fetchall()
    by_path = {
        row["relative_path"]: row
        for row in con.execute(
            "SELECT * FROM source_record WHERE status='present'"
        ).fetchall()
    }
    result: Dict[str, Dict[str, Any]] = {}
    for archive_row in rows:
        metadata = _row_metadata(archive_row)
        if metadata.get("archive_semantic_status") != "duplicate_of_extracted":
            continue
        members = metadata.get("archive_members")
        if not isinstance(members, list):
            continue
        for member in members:
            if not isinstance(member, dict) or member.get("is_directory"):
                continue
            member_path = str(member.get("member_path") or "")
            directory_path = _archive_directory_path(
                config,
                archive_row,
                member_path,
            )
            directory_row = by_path.get(directory_path or "")
            if directory_row is None or not directory_row["source_version_id"]:
                continue
            evidence_rows = con.execute(
                """
                SELECT evidence_id, locator
                FROM evidence_record
                WHERE source_version_id=?
                ORDER BY locator
                """,
                (directory_row["source_version_id"],),
            ).fetchall()
            semantic_evidence_ids = [row["evidence_id"] for row in evidence_rows]
            semantic_locators = [row["locator"] for row in evidence_rows]
            if semantic_evidence_ids:
                semantic_id = semantic_evidence_ids[0]
            else:
                semantic_id = stable_evidence_id(
                    directory_row["source_version_id"],
                    "document",
                )
            item = result.setdefault(
                semantic_id,
                {
                    "semantic_evidence_id": semantic_id,
                    "semantic_evidence_ids": semantic_evidence_ids or [semantic_id],
                    "semantic_locators": semantic_locators or ["document"],
                    "directory_relative_path": directory_path,
                    "directory_source_id": directory_row["source_id"],
                    "source_references": [],
                    "archive_references": [],
                    "semantic_status": "duplicate_of_extracted",
                },
            )
            references = set(item["source_references"])
            references.update(
                {directory_row["source_id"], archive_row["source_id"]}
            )
            item["source_references"] = sorted(references)
            item["archive_references"].append(
                {
                    "archive_source_id": archive_row["source_id"],
                    "archive_path": archive_row["relative_path"],
                    "member_path": member_path,
                }
            )
    return result


def _upsert_normalized_status(
    con: sqlite3.Connection,
    row: sqlite3.Row,
    *,
    status: str,
    error: Optional[str],
    normalized_path: Optional[str] = None,
    parser: Optional[str] = None,
    parser_version: Optional[str] = None,
    char_count: Optional[int] = None,
    line_count: Optional[int] = None,
    content_sha256: Optional[str] = None,
) -> None:
    timestamp = now_iso()
    con.execute(
        """
        INSERT INTO normalized_document (
            source_version_id, source_id, derived_path, normalized_path,
            parser, parser_version, source_mtime_ns, char_count, line_count,
            content_sha256, status, error, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(source_version_id) DO UPDATE SET
            source_id=excluded.source_id,
            source_mtime_ns=excluded.source_mtime_ns,
            status=excluded.status,
            error=excluded.error,
            updated_at=excluded.updated_at,
            derived_path=COALESCE(excluded.derived_path, normalized_document.derived_path),
            normalized_path=COALESCE(excluded.normalized_path, normalized_document.normalized_path),
            parser=COALESCE(excluded.parser, normalized_document.parser),
            parser_version=COALESCE(excluded.parser_version, normalized_document.parser_version),
            char_count=COALESCE(excluded.char_count, normalized_document.char_count),
            line_count=COALESCE(excluded.line_count, normalized_document.line_count),
            content_sha256=COALESCE(excluded.content_sha256, normalized_document.content_sha256)
        """,
        (
            row["source_version_id"],
            row["source_id"],
            normalized_path,
            normalized_path,
            parser,
            parser_version,
            row["mtime_ns"],
            char_count,
            line_count,
            content_sha256,
            status,
            error,
            timestamp,
        ),
    )


def _derived_text(
    config: Config,
    row: sqlite3.Row,
    extracted: ExtractionResult,
) -> Tuple[str, str]:
    metadata = _row_metadata(row)
    extracted_metadata = dict(extracted.metadata)
    date_basis = (
        extracted_metadata.get("source_date_basis")
        or metadata.get("source_date_basis")
        or ("filename_hint" if row["date_hint"] else "not_observed")
    )
    locator = extracted.parts[0].locator if len(extracted.parts) == 1 else "multi"
    header = provenance_header(
        {
            "source_id": row["source_id"],
            "source_version_id": row["source_version_id"],
            "source_system": row["source_system"],
            "original_path": row["absolute_path"],
            "relative_path": row["relative_path"],
            "original_date": extracted_metadata.get("event_date") or row["date_hint"] or "",
            "source_date_basis": date_basis,
            "file_type": row["extension"] or row["kind"],
            "parser": extracted.parser,
            "parser_version": extracted.parser_version,
            "locator": locator,
            "generated_at": now_iso(),
        }
    )
    sections = [
        f"## Locator: {part.locator}\n\n{part.text.rstrip()}"
        for part in extracted.parts
        if part.text.strip()
    ]
    return scrub_derived_text(header + "\n\n".join(sections) + "\n"), date_basis


def _record_extraction_review(
    con: sqlite3.Connection,
    row: sqlite3.Row,
    issue_type: str,
    error: str,
) -> None:
    record_review_item(
        con,
        issue_type=issue_type,
        source_id=row["source_id"],
        proposed_result={
            "source_version_id": row["source_version_id"],
            "status": "review_required",
            "error": scrub_derived_text(error),
        },
        reason=scrub_derived_text(error),
    )


def _resolve_extraction_reviews(
    con: sqlite3.Connection,
    row: sqlite3.Row,
    parser: str,
) -> None:
    """Resolve a prior parser review only when the same source version succeeds."""
    source_version_id = str(row["source_version_id"] or "")
    if not source_version_id:
        return
    candidates = con.execute(
        """
        SELECT review_id, proposed_result_json
        FROM review_item
        WHERE source_id=? AND status='pending'
          AND issue_type IN ('parser_unsupported', 'parser_error', 'extraction_blocked')
        """,
        (row["source_id"],),
    ).fetchall()
    for review in candidates:
        try:
            proposed = json.loads(review["proposed_result_json"] or "{}")
        except (TypeError, ValueError, json.JSONDecodeError):
            proposed = {}
        if not isinstance(proposed, dict) or str(proposed.get("source_version_id") or "") != source_version_id:
            continue
        con.execute(
            """
            UPDATE review_item
            SET status='resolved',
                resolution=?,
                updated_at=?
            WHERE review_id=? AND status='pending'
            """,
            (
                f"Superseded: source version normalized successfully with {parser}.",
                now_iso(),
                review["review_id"],
            ),
        )


def _task_date_basis(source_date_basis: str) -> str:
    """Map inventory provenance names to the task date vocabulary."""
    return {
        "meeting_folder": "meeting_date",
        "meeting": "meeting_date",
        "event": "event_date",
        "event_date": "event_date",
    }.get(source_date_basis, source_date_basis)


def _record_date_observation(
    con: sqlite3.Connection,
    *,
    source_version_id: str,
    evidence_id: str,
    date_value: str,
    basis: str,
) -> None:
    if not date_value or not basis:
        return
    observation_id = stable_id(
        "date-observation",
        source_version_id,
        evidence_id,
        date_value,
        basis,
    )
    con.execute(
        """
        INSERT OR IGNORE INTO date_observation (
            observation_id, evidence_id, source_version_id, date_value,
            basis, precision, confidence, created_at
        ) VALUES (?, ?, ?, ?, ?, 'day', 0.95, ?)
        """,
        (
            observation_id,
            evidence_id,
            source_version_id,
            date_value,
            basis,
            now_iso(),
        ),
    )


def normalize_all(config: Config, con: sqlite3.Connection) -> Dict[str, int]:
    include_unclassified_sources = bool(
        config.get(
            "normalization",
            "include_unclassified_sources",
            False,
        )
    )
    skip_classifications = {
        str(value).casefold()
        for value in config.get(
            "normalization",
            "skip_classifications",
            list(DEFAULT_SKIP_CLASSIFICATIONS),
        )
        if str(value).strip()
    }
    if include_unclassified_sources:
        # The caller has explicitly asked for a complete local corpus. Keep
        # every parseable source, regardless of the inventory's provisional
        # scope label or whether its path matched a configured root.
        skip_classifications = set()

    mark_noncurrent_normalized_documents_retained(
        con,
        now_iso(),
        skip_classifications=skip_classifications,
    )

    eligible = con.execute(
        """
        SELECT s.*
        FROM source_record s
        WHERE s.status = 'present'
          AND s.extraction_status = 'ready'
          AND s.content_sha256 IS NOT NULL
          AND s.content_sha256 <> ''
          AND s.source_version_id IS NOT NULL
          AND s.source_version_id <> ''
          AND EXISTS (
              SELECT 1
              FROM source_version v
              WHERE v.source_version_id=s.source_version_id
                AND v.source_id=s.source_id
                AND v.content_sha256=s.content_sha256
          )
          AND s.kind NOT IN (
            'media', 'email_data', 'mcp', 'unknown',
              'discovery', 'artifact'
          )
        ORDER BY s.source_system, s.relative_path
        """
    ).fetchall()
    if include_unclassified_sources:
        approved_root = eligible
        unapproved_root = []
    else:
        approved_root = [row for row in eligible if _has_explicit_source_root(row)]
        unapproved_root = [row for row in eligible if not _has_explicit_source_root(row)]
    rows = [
        row
        for row in approved_root
        if (row["classification"] or "Unknown").casefold() not in skip_classifications
    ]
    if include_unclassified_sources:
        review_rows = [
            row
            for row in eligible
            if (row["classification"] or "Unknown").casefold() in skip_classifications
        ]
    else:
        review_rows = [
            row
            for row in approved_root + unapproved_root
            if (
                not _has_explicit_source_root(row)
                or (row["classification"] or "Unknown").casefold() in skip_classifications
            )
        ]

    result = {
        "normalized": 0,
        "skipped_unchanged": 0,
        "review_required": len(review_rows),
        "unsupported": 0,
        "errors": 0,
    }

    for row in review_rows:
        error = (
            f"Skipped pending source review: classification={row['classification']}"
            + (
                "; source root is not explicitly approved"
                if not _has_explicit_source_root(row)
                else ""
            )
        )
        _upsert_normalized_status(
            con,
            row,
            status="review_required",
            error=error,
        )
        record_scope_proposal(
            con,
            row["source_id"],
            "Unknown" if not _has_explicit_source_root(row) else (row["classification"] or "Unknown"),
            sensitivity=row["sensitivity"] or "unknown",
            reason=error,
        )

    for row in rows:
        source_version_id = row["source_version_id"]
        existing = con.execute(
            """
            SELECT status, normalized_path, parser_version
            FROM normalized_document
            WHERE source_version_id = ?
            """,
            (source_version_id,),
        ).fetchone()
        expected_parser_version = (
            EMAIL_PARSER_VERSION
            if str(row["extension"] or "").casefold() == ".eml"
            else PARSER_VERSION
        )
        if (
            existing
            and existing["status"] in {"normalized", "metadata", "prior_good_retained"}
            and existing["parser_version"] == expected_parser_version
            and existing["normalized_path"]
            and Path(existing["normalized_path"]).exists()
        ):
            result["skipped_unchanged"] += 1
            continue

        path = Path(row["absolute_path"])
        output = _output_path(config, row["source_system"], source_version_id)
        try:
            extracted = extract_source(
                path,
                row["kind"],
                row["extension"],
                config,
                source_system=row["source_system"],
                relative_path=row["relative_path"],
            )
            text, _date_basis = _derived_text(config, row, extracted)
            atomic_write_text(output, text)
            _upsert_normalized_status(
                con,
                row,
                status=extracted.status,
                error=None,
                normalized_path=str(output),
                parser=extracted.parser,
                parser_version=extracted.parser_version,
                char_count=len(text),
                line_count=text.count("\n") + 1,
                content_sha256=sha256_text(text),
            )
            date_recorded = False
            for part in extracted.parts:
                part_text = scrub_derived_text(part.text).strip()
                evidence_id = record_evidence(
                    con,
                    source_version_id=source_version_id,
                    locator=part.locator,
                    derived_text_path=str(output),
                    text_sha256=sha256_text(part_text),
                    evidence_status=(
                        "metadata" if extracted.status == "metadata" else "derived"
                    ),
                    derived_text=(
                        None
                        if extracted.status == "metadata"
                        else part_text
                    ),
                )
                if extracted.status == "metadata":
                    con.execute(
                        "DELETE FROM derived_text_fts WHERE evidence_id=?",
                        (evidence_id,),
                    )
                    continue
                if not date_recorded and extracted.metadata.get("event_date"):
                    _record_date_observation(
                        con,
                        source_version_id=source_version_id,
                        evidence_id=evidence_id,
                        date_value=str(extracted.metadata["event_date"]),
                        basis=str(
                            extracted.metadata.get("source_date_basis")
                            or "event_date"
                        ),
                    )
                    date_recorded = True
                extract_task_proposals(
                    con,
                    part_text,
                    source_id=row["source_id"],
                    source_evidence_id=evidence_id,
                    source_event_date=(
                        extracted.metadata.get("event_date") or row["date_hint"]
                    ),
                    source_date_basis=_task_date_basis(_date_basis),
                    run_date=date.today(),
                    scope=row["classification"] or "Unknown",
                    commit=False,
                )
            _resolve_extraction_reviews(con, row, extracted.parser)
            result["normalized"] += 1
        except ExtractionBlocked as exc:
            error = scrub_derived_text(str(exc))
            _upsert_normalized_status(
                con,
                row,
                status="blocked",
                error=error,
            )
            _record_extraction_review(
                con,
                row,
                "extraction_blocked",
                error,
            )
            result["unsupported"] += 1
        except RuntimeError as exc:
            error = scrub_derived_text(str(exc))
            _upsert_normalized_status(
                con,
                row,
                status="unsupported",
                error=error,
            )
            _record_extraction_review(
                con,
                row,
                "parser_unsupported",
                error,
            )
            result["unsupported"] += 1
        except Exception as exc:
            error = scrub_derived_text(f"{type(exc).__name__}: {exc}")
            _upsert_normalized_status(
                con,
                row,
                status="error",
                error=error,
            )
            _record_extraction_review(con, row, "parser_error", error)
            result["errors"] += 1
    con.commit()
    archive_map_path = config.state_dir / "archive_reference_map.json"
    config.assert_derived_path(archive_map_path)
    atomic_write_json(
        archive_map_path,
        build_archive_reference_map(config, con),
    )
    return result
